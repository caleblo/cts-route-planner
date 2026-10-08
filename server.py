#!/usr/bin/env python3
"""
CTS Route Planner — backend (Python stdlib + openpyxl).

Serves a Leaflet front end for pin-dropping / location saving, and exposes
a small JSON API that:
  - geocodes place names   (Nominatim / OpenStreetMap)
  - reverse-geocodes pins  (Nominatim)
  - builds a distance+duration matrix (OSRM /table)
  - computes an optimized stop order (OSRM /trip)
  - gets the real road geometry for the ordered route (OSRM /route)
  - calculates fuel cost total + per-stop breakdown
  - adds per-stop service (handling) time + checks the 7-hour shift limit
  - writes route documents (Markdown + CSV[utf-8-sig] + JSON + XLSX) to ./routes/
  - saves / loads whole sessions (stops, nicknames, rates, vehicle) to ./sessions/

No API key required. Only non-stdlib dependency is openpyxl (for .xlsx).

Run:  python server.py            (default http://127.0.0.1:8777)
      python server.py --port 9000
"""

import argparse
import csv
import io
import json
import os
import re
import time
import urllib.parse
import urllib.request
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    HAVE_XLSX = True
except Exception:
    HAVE_XLSX = False

# --------------------------------------------------------------------------- #
# Config
# --------------------------------------------------------------------------- #
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
STATIC_DIR = os.path.join(BASE_DIR, "static")
ROUTES_DIR = os.path.join(BASE_DIR, "routes")
SESSIONS_DIR = os.path.join(BASE_DIR, "sessions")
os.makedirs(ROUTES_DIR, exist_ok=True)
os.makedirs(SESSIONS_DIR, exist_ok=True)

NOMINATIM = "https://nominatim.openstreetmap.org"
OSRM = "https://router.project-osrm.org"
USER_AGENT = "CTS-RoutePlanner/2.0 (self-hosted tool)"

# Shift / delivery-time policy (default; overridable per request)
MAX_SHIFT_HOURS = 7.0

# Vehicle presets. economy in L/100km, service_min per stop,
# speed_kmh = average road speed used to estimate driving time
# (0 / None = use OSRM's own duration estimate).
VEHICLES = {
    "truck_5t": {"label": "5-tonne truck", "economy": 20.0, "unit": "l_per_100km",
                 "service_min": 15.0, "speed_kmh": 35.0},
    "van": {"label": "Delivery van", "economy": 11.0, "unit": "l_per_100km",
            "service_min": 10.0, "speed_kmh": 40.0},
    "car": {"label": "Car", "economy": 8.0, "unit": "l_per_100km",
            "service_min": 5.0, "speed_kmh": 0.0},
}

_last_nominatim_call = [0.0]  # throttle: Nominatim ToS = max 1 req/s


# --------------------------------------------------------------------------- #
# HTTP helpers
# --------------------------------------------------------------------------- #
def _get_json(url, throttle_nominatim=False):
    if throttle_nominatim:
        elapsed = time.time() - _last_nominatim_call[0]
        if elapsed < 1.1:
            time.sleep(1.1 - elapsed)
        _last_nominatim_call[0] = time.time()
    req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(req, timeout=30) as resp:
        return json.loads(resp.read().decode("utf-8"))


# --------------------------------------------------------------------------- #
# Geo services
# --------------------------------------------------------------------------- #
def geocode(query):
    q = urllib.parse.quote(query)
    url = f"{NOMINATIM}/search?q={q}&format=json&limit=5&addressdetails=1"
    data = _get_json(url, throttle_nominatim=True)
    out = []
    for r in data:
        out.append({
            "name": r.get("display_name", query).split(",")[0],
            "display_name": r.get("display_name", ""),
            "lat": float(r["lat"]),
            "lon": float(r["lon"]),
        })
    return out


def reverse_geocode(lat, lon):
    url = f"{NOMINATIM}/reverse?lat={lat}&lon={lon}&format=json&addressdetails=1"
    try:
        data = _get_json(url, throttle_nominatim=True)
    except Exception:
        return {"name": f"{lat:.5f}, {lon:.5f}", "display_name": ""}
    addr = data.get("address", {})
    name = (addr.get("building") or addr.get("amenity") or addr.get("road")
            or data.get("name") or data.get("display_name", "").split(",")[0]
            or f"{lat:.5f}, {lon:.5f}")
    return {"name": name, "display_name": data.get("display_name", "")}


def osrm_table(coords):
    coord_str = ";".join(f"{lon},{lat}" for lat, lon in coords)
    url = f"{OSRM}/table/v1/driving/{coord_str}?annotations=duration,distance"
    data = _get_json(url)
    if data.get("code") != "Ok":
        raise RuntimeError(f"OSRM table error: {data.get('code')}")
    return data["distances"], data["durations"]


def osrm_trip(coords, roundtrip=False):
    coord_str = ";".join(f"{lon},{lat}" for lat, lon in coords)
    rt = "true" if roundtrip else "false"
    url = (f"{OSRM}/trip/v1/driving/{coord_str}"
           f"?source=first&roundtrip={rt}&overview=false")
    data = _get_json(url)
    if data.get("code") != "Ok":
        raise RuntimeError(f"OSRM trip error: {data.get('code')}")
    order = [None] * len(coords)
    for i, wp in enumerate(data["waypoints"]):
        order[wp["waypoint_index"]] = i
    return [o for o in order if o is not None]


def osrm_route(coords):
    coord_str = ";".join(f"{lon},{lat}" for lat, lon in coords)
    url = (f"{OSRM}/route/v1/driving/{coord_str}"
           f"?overview=full&geometries=geojson&annotations=false")
    data = _get_json(url)
    if data.get("code") != "Ok":
        raise RuntimeError(f"OSRM route error: {data.get('code')}")
    route = data["routes"][0]
    geometry = route["geometry"]["coordinates"]
    legs = [{"distance": leg["distance"], "duration": leg["duration"]}
            for leg in route["legs"]]
    return {
        "geometry": [[lat, lon] for lon, lat in geometry],
        "legs": legs,
        "total_distance": route["distance"],
        "total_duration": route["duration"],
    }


# --------------------------------------------------------------------------- #
# Core: build a full route plan with fuel cost + per-stop breakdown + timing
# --------------------------------------------------------------------------- #
def _stop_label(s):
    """Prefer nickname, fall back to name."""
    nick = (s.get("nickname") or "").strip()
    return nick if nick else s.get("name", "Stop")


def build_route_plan(stops, fuel, vehicle=None, optimize=True, roundtrip=False,
                     max_hours=None, speed_kmh=None):
    if len(stops) < 2:
        raise ValueError("Need at least 2 stops.")

    max_hours = float(max_hours) if max_hours else MAX_SHIFT_HOURS
    # Average speed override (km/h). <=0 or None => use OSRM's own durations.
    try:
        speed_kmh = float(speed_kmh) if speed_kmh else 0.0
    except (TypeError, ValueError):
        speed_kmh = 0.0
    m_per_s = speed_kmh * 1000.0 / 3600.0 if speed_kmh > 0 else 0.0

    coords = [(s["lat"], s["lon"]) for s in stops]
    distances, durations = osrm_table(coords)

    if optimize and len(stops) >= 3:
        order = osrm_trip(coords, roundtrip=roundtrip)
    else:
        order = list(range(len(stops)))

    ordered_stops = [stops[i] for i in order]
    ordered_coords = [(s["lat"], s["lon"]) for s in ordered_stops]
    if roundtrip:
        ordered_coords.append(ordered_coords[0])

    route = osrm_route(ordered_coords)

    # Fuel economy -> litres per metre
    price = float(fuel.get("price", 0))
    economy = float(fuel.get("economy", 0))
    unit = fuel.get("unit", "l_per_100km")
    if unit == "km_per_l":
        l_per_m = (1.0 / economy) / 1000.0 if economy else 0.0
    else:
        l_per_m = (economy / 100.0) / 1000.0

    legs = route["legs"]
    seq_stops = ordered_stops + ([ordered_stops[0]] if roundtrip else [])

    breakdown = []
    cum_dist = 0.0
    cum_drive = 0.0       # driving seconds
    cum_service = 0.0     # service seconds
    cum_cost = 0.0

    # First stop = origin (service time still applies at the depot start)
    first = ordered_stops[0]
    first_service = float(first.get("service_min", 0)) * 60.0
    cum_service += first_service
    breakdown.append({
        "seq": 1,
        "name": _stop_label(first),
        "full_name": first.get("name", ""),
        "lat": first["lat"],
        "lon": first["lon"],
        "leg_from": "— (start)",
        "leg_distance_km": 0.0,
        "leg_drive_min": 0.0,
        "service_min": round(first_service / 60.0, 1),
        "leg_fuel_l": 0.0,
        "leg_cost": 0.0,
        "cum_distance_km": 0.0,
        "cum_time_min": round(cum_service / 60.0, 1),
        "cum_cost": 0.0,
    })

    for idx, leg in enumerate(legs):
        dest = seq_stops[idx + 1]
        src = seq_stops[idx]
        leg_dist = leg["distance"]
        # Drive time: from the chosen average speed, else OSRM's estimate.
        leg_drive = (leg_dist / m_per_s) if m_per_s > 0 else leg["duration"]
        leg_litres = leg_dist * l_per_m
        leg_cost = leg_litres * price
        svc = float(dest.get("service_min", 0)) * 60.0
        cum_dist += leg_dist
        cum_drive += leg_drive
        cum_service += svc
        cum_cost += leg_cost
        breakdown.append({
            "seq": idx + 2,
            "name": _stop_label(dest),
            "full_name": dest.get("name", ""),
            "lat": dest["lat"],
            "lon": dest["lon"],
            "leg_from": _stop_label(src),
            "leg_distance_km": round(leg_dist / 1000.0, 3),
            "leg_drive_min": round(leg_drive / 60.0, 1),
            "service_min": round(svc / 60.0, 1),
            "leg_fuel_l": round(leg_litres, 3),
            "leg_cost": round(leg_cost, 2),
            "cum_distance_km": round(cum_dist / 1000.0, 3),
            "cum_time_min": round((cum_drive + cum_service) / 60.0, 1),
            "cum_cost": round(cum_cost, 2),
        })

    total_dist_km = route["total_distance"] / 1000.0
    total_drive_min = ((route["total_distance"] / m_per_s) if m_per_s > 0
                       else route["total_duration"]) / 60.0
    total_service_min = cum_service / 60.0
    total_delivery_min = total_drive_min + total_service_min
    total_litres = route["total_distance"] * l_per_m
    total_cost = total_litres * price

    within_limit = total_delivery_min <= max_hours * 60.0

    plan = {
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "optimize": optimize,
        "roundtrip": roundtrip,
        "vehicle": vehicle or {},
        "fuel": {"price": price, "economy": economy, "unit": unit,
                 "l_per_m": l_per_m, "speed_kmh": speed_kmh},
        "order_indices": order,
        "ordered_stops": [{
            "name": s.get("name", ""),
            "nickname": s.get("nickname", ""),
            "lat": s["lat"], "lon": s["lon"],
            "service_min": float(s.get("service_min", 0)),
        } for s in ordered_stops],
        "matrix": {
            "labels": [_stop_label(s) for s in stops],
            "distances_km": [[round(d / 1000.0, 3) for d in row]
                             for row in distances],
            "durations_min": [[round(d / 60.0, 1) for d in row]
                              for row in durations],
        },
        "breakdown": breakdown,
        "geometry": route["geometry"],
        "totals": {
            "distance_km": round(total_dist_km, 3),
            "drive_min": round(total_drive_min, 1),
            "service_min": round(total_service_min, 1),
            "delivery_min": round(total_delivery_min, 1),
            "delivery_hours": round(total_delivery_min / 60.0, 2),
            "fuel_l": round(total_litres, 3),
            "cost": round(total_cost, 2),
            "num_stops": len(ordered_stops),
        },
        "shift": {
            "max_hours": max_hours,
            "within_limit": within_limit,
            "over_by_min": round(max(0.0, total_delivery_min
                                     - max_hours * 60.0), 1),
        },
    }
    return plan


# --------------------------------------------------------------------------- #
# Document writers
# --------------------------------------------------------------------------- #
def _fmt_hm(minutes):
    h = int(minutes // 60)
    m = int(round(minutes - h * 60))
    return f"{h}h {m:02d}m"


def write_documents(plan):
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    stem = f"route_{ts}"
    paths = {}
    t = plan["totals"]
    f = plan["fuel"]
    sh = plan["shift"]
    veh = plan.get("vehicle", {})
    unit_label = ("L/100km" if f["unit"] == "l_per_100km" else "km/L")
    shift_msg = (f"WITHIN {sh['max_hours']:g}-hour limit" if sh["within_limit"]
                 else f"OVER {sh['max_hours']:g}-hour limit by "
                      f"{_fmt_hm(sh['over_by_min'])}")

    # --- Markdown ---
    md = io.StringIO()
    md.write(f"# Route Plan — {plan['generated_at']}\n\n## Summary\n\n")
    md.write(f"- **Stops:** {t['num_stops']}\n")
    md.write(f"- **Optimized order:** {'yes' if plan['optimize'] else 'no'}"
             f"{' (round trip)' if plan['roundtrip'] else ''}\n")
    md.write(f"- **Total distance:** {t['distance_km']:.2f} km\n")
    spd = f.get("speed_kmh", 0)
    if spd and spd > 0:
        md.write(f"- **Driving time:** {_fmt_hm(t['drive_min'])} "
                 f"(avg speed {spd:.0f} km/h)\n")
    else:
        md.write(f"- **Driving time:** {_fmt_hm(t['drive_min'])} "
                 f"(OSRM estimate)\n")
    md.write(f"- **Service time:** {_fmt_hm(t['service_min'])}\n")
    md.write(f"- **TOTAL DELIVERY TIME:** **{_fmt_hm(t['delivery_min'])}** "
             f"({t['delivery_hours']:.2f} h)\n")
    md.write(f"- **{sh['max_hours']:g}-hour check:** **{shift_msg}**\n")
    md.write(f"- **Fuel price:** {f['price']:.2f} / L\n")
    md.write(f"- **Fuel economy:** {f['economy']:.2f} {unit_label}\n")
    md.write(f"- **Total fuel:** {t['fuel_l']:.2f} L\n")
    md.write(f"- **TOTAL FUEL COST:** **{t['cost']:.2f}**\n\n")

    md.write("## Optimized stop order\n\n")
    for b in plan["breakdown"]:
        md.write(f"{b['seq']}. {b['name']}  ({b['lat']:.6f}, {b['lon']:.6f})\n")
    md.write("\n## Per-stop breakdown\n\n")
    md.write("| # | Stop | From | Leg km | Drive min | Service min | Leg L | "
             "Leg cost | Cum km | Cum time | Cum cost |\n")
    md.write("|---|------|------|-------:|----------:|------------:|------:|"
             "---------:|-------:|---------:|---------:|\n")
    for b in plan["breakdown"]:
        md.write(f"| {b['seq']} | {b['name']} | {b['leg_from']} | "
                 f"{b['leg_distance_km']:.2f} | {b['leg_drive_min']:.1f} | "
                 f"{b['service_min']:.1f} | {b['leg_fuel_l']:.2f} | "
                 f"{b['leg_cost']:.2f} | {b['cum_distance_km']:.2f} | "
                 f"{_fmt_hm(b['cum_time_min'])} | {b['cum_cost']:.2f} |\n")
    md.write(f"| | **TOTAL** | | **{t['distance_km']:.2f}** | "
             f"**{t['drive_min']:.1f}** | **{t['service_min']:.1f}** | "
             f"**{t['fuel_l']:.2f}** | **{t['cost']:.2f}** | | "
             f"**{_fmt_hm(t['delivery_min'])}** | **{t['cost']:.2f}** |\n\n")

    md.write("## Distance matrix (km)\n\n")
    labels = plan["matrix"]["labels"]
    md.write("| From \\ To | " + " | ".join(labels) + " |\n")
    md.write("|" + "---|" * (len(labels) + 1) + "\n")
    for i, row in enumerate(plan["matrix"]["distances_km"]):
        md.write(f"| {labels[i]} | " + " | ".join(f"{v:.2f}" for v in row)
                 + " |\n")
    md_path = os.path.join(ROUTES_DIR, stem + ".md")
    with open(md_path, "w", encoding="utf-8") as fh:
        fh.write(md.getvalue())
    paths["markdown"] = md_path

    # --- CSV (breakdown) — utf-8-sig so Excel reads HK/Chinese text correctly ---
    csv_path = os.path.join(ROUTES_DIR, stem + "_breakdown.csv")
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(["seq", "stop", "full_name", "from", "leg_distance_km",
                    "leg_drive_min", "service_min", "leg_fuel_l", "leg_cost",
                    "cum_distance_km", "cum_time_min", "cum_cost", "lat", "lon"])
        for b in plan["breakdown"]:
            w.writerow([b["seq"], b["name"], b["full_name"], b["leg_from"],
                        b["leg_distance_km"], b["leg_drive_min"],
                        b["service_min"], b["leg_fuel_l"], b["leg_cost"],
                        b["cum_distance_km"], b["cum_time_min"], b["cum_cost"],
                        b["lat"], b["lon"]])
    paths["csv"] = csv_path

    # --- Matrix CSV (utf-8-sig) ---
    matrix_csv = os.path.join(ROUTES_DIR, stem + "_matrix.csv")
    with open(matrix_csv, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(["distance_km / to ->"] + labels)
        for i, row in enumerate(plan["matrix"]["distances_km"]):
            w.writerow([labels[i]] + row)
        w.writerow([])
        w.writerow(["duration_min / to ->"] + labels)
        for i, row in enumerate(plan["matrix"]["durations_min"]):
            w.writerow([labels[i]] + row)
    paths["matrix_csv"] = matrix_csv

    # --- JSON ---
    json_path = os.path.join(ROUTES_DIR, stem + ".json")
    with open(json_path, "w", encoding="utf-8") as fh:
        json.dump(plan, fh, indent=2, ensure_ascii=False)
    paths["json"] = json_path

    # --- XLSX ---
    if HAVE_XLSX:
        xlsx_path = os.path.join(ROUTES_DIR, stem + ".xlsx")
        _write_xlsx(xlsx_path, plan)
        paths["xlsx"] = xlsx_path

    return paths


def _write_xlsx(path, plan):
    t = plan["totals"]
    f = plan["fuel"]
    sh = plan["shift"]
    veh = plan.get("vehicle", {})
    unit_label = ("L/100km" if f["unit"] == "l_per_100km" else "km/L")

    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    hdr_font = Font(bold=True, color="FFFFFF")
    tot_fill = PatternFill("solid", fgColor="DDEBF7")
    bold = Font(bold=True)
    ok_fill = PatternFill("solid", fgColor="C6EFCE")
    bad_fill = PatternFill("solid", fgColor="FFC7CE")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center")

    wb = Workbook()

    # Sheet 1: Summary
    ws = wb.active
    ws.title = "Summary"
    rows = [
        ("Route Plan", plan["generated_at"]),
        ("Stops", t["num_stops"]),
        ("Optimized order", "yes" if plan["optimize"] else "no"),
        ("Round trip", "yes" if plan["roundtrip"] else "no"),
        ("Total distance (km)", t["distance_km"]),
        ("Avg speed (km/h)", (f"{f.get('speed_kmh', 0):.0f}"
                              if f.get("speed_kmh", 0) else "OSRM estimate")),
        ("Driving time", _fmt_hm(t["drive_min"])),
        ("Service time", _fmt_hm(t["service_min"])),
        ("TOTAL DELIVERY TIME", f"{_fmt_hm(t['delivery_min'])} "
                                f"({t['delivery_hours']:.2f} h)"),
        (f"Within {sh['max_hours']:g}-hour limit?",
         "YES" if sh["within_limit"] else
         f"NO — over by {_fmt_hm(sh['over_by_min'])}"),
        ("Fuel price / L", f["price"]),
        ("Fuel economy", f"{f['economy']:.2f} {unit_label}"),
        ("Total fuel (L)", t["fuel_l"]),
        ("TOTAL FUEL COST", t["cost"]),
    ]
    for r, (k, v) in enumerate(rows, 1):
        ws.cell(r, 1, k).font = bold
        c = ws.cell(r, 2, v)
        if k.startswith("TOTAL") or k.startswith("Within"):
            ws.cell(r, 1).font = Font(bold=True, size=12)
            c.font = Font(bold=True, size=12)
        if k.startswith("Within"):
            c.fill = ok_fill if sh["within_limit"] else bad_fill
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 34

    # Sheet 2: Breakdown
    ws2 = wb.create_sheet("Per-stop breakdown")
    cols = ["#", "Stop", "From", "Leg km", "Drive min", "Service min",
            "Leg L", "Leg cost", "Cum km", "Cum time", "Cum cost"]
    for ci, name in enumerate(cols, 1):
        c = ws2.cell(1, ci, name)
        c.fill = hdr_fill; c.font = hdr_font; c.alignment = center; c.border = border
    for ri, b in enumerate(plan["breakdown"], 2):
        vals = [b["seq"], b["name"], b["leg_from"], b["leg_distance_km"],
                b["leg_drive_min"], b["service_min"], b["leg_fuel_l"],
                b["leg_cost"], b["cum_distance_km"], _fmt_hm(b["cum_time_min"]),
                b["cum_cost"]]
        for ci, v in enumerate(vals, 1):
            cell = ws2.cell(ri, ci, v); cell.border = border
    rtot = len(plan["breakdown"]) + 2
    tot_vals = ["", "TOTAL", "", t["distance_km"], t["drive_min"],
                t["service_min"], t["fuel_l"], t["cost"], t["distance_km"],
                _fmt_hm(t["delivery_min"]), t["cost"]]
    for ci, v in enumerate(tot_vals, 1):
        cell = ws2.cell(rtot, ci, v); cell.font = bold; cell.fill = tot_fill
        cell.border = border
    widths2 = [5, 30, 30, 9, 11, 12, 8, 10, 9, 11, 10]
    for i, w in enumerate(widths2, 1):
        ws2.column_dimensions[chr(64 + i)].width = w
    ws2.freeze_panes = "A2"

    # Sheet 3: Distance matrix
    ws3 = wb.create_sheet("Distance matrix km")
    labels = plan["matrix"]["labels"]
    c = ws3.cell(1, 1, "From \\ To"); c.fill = hdr_fill; c.font = hdr_font
    c.border = border
    for ci, lab in enumerate(labels, 2):
        c = ws3.cell(1, ci, lab); c.fill = hdr_fill; c.font = hdr_font
        c.alignment = center; c.border = border
    for ri, row in enumerate(plan["matrix"]["distances_km"], 2):
        c = ws3.cell(ri, 1, labels[ri - 2]); c.font = bold; c.border = border
        for ci, v in enumerate(row, 2):
            cc = ws3.cell(ri, ci, v); cc.border = border
    ws3.column_dimensions["A"].width = 26
    for i in range(len(labels)):
        ws3.column_dimensions[chr(66 + i) if i < 25 else "AA"].width = 14
    ws3.freeze_panes = "B2"

    # Sheet 4: Duration matrix
    ws4 = wb.create_sheet("Duration matrix min")
    c = ws4.cell(1, 1, "From \\ To"); c.fill = hdr_fill; c.font = hdr_font
    c.border = border
    for ci, lab in enumerate(labels, 2):
        c = ws4.cell(1, ci, lab); c.fill = hdr_fill; c.font = hdr_font
        c.alignment = center; c.border = border
    for ri, row in enumerate(plan["matrix"]["durations_min"], 2):
        c = ws4.cell(ri, 1, labels[ri - 2]); c.font = bold; c.border = border
        for ci, v in enumerate(row, 2):
            cc = ws4.cell(ri, ci, v); cc.border = border
    ws4.column_dimensions["A"].width = 26
    ws4.freeze_panes = "B2"

    wb.save(path)


# --------------------------------------------------------------------------- #
# Sessions
# --------------------------------------------------------------------------- #
_SAFE = re.compile(r"[^A-Za-z0-9_\-]+")


def _safe_name(name):
    name = _SAFE.sub("_", (name or "").strip()) or "session"
    return name[:64]


def save_session(name, payload):
    fn = _safe_name(name) + ".json"
    path = os.path.join(SESSIONS_DIR, fn)
    record = {"name": name, "saved_at": datetime.now().isoformat(timespec="seconds"),
              "data": payload}
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(record, fh, indent=2, ensure_ascii=False)
    return {"saved": fn, "path": path}


def list_sessions():
    out = []
    for fn in sorted(os.listdir(SESSIONS_DIR)):
        if not fn.endswith(".json"):
            continue
        try:
            with open(os.path.join(SESSIONS_DIR, fn), encoding="utf-8") as fh:
                rec = json.load(fh)
            out.append({"file": fn, "name": rec.get("name", fn[:-5]),
                        "saved_at": rec.get("saved_at", ""),
                        "stops": len(rec.get("data", {}).get("stops", []))})
        except Exception:
            continue
    return out


def load_session(fn):
    path = os.path.join(SESSIONS_DIR, _safe_name(fn.replace(".json", "")) + ".json")
    if not os.path.isfile(path):
        raise FileNotFoundError(fn)
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


# --------------------------------------------------------------------------- #
# HTTP handler
# --------------------------------------------------------------------------- #
class Handler(BaseHTTPRequestHandler):
    def log_message(self, fmt, *args):
        print(f"[{self.log_date_time_string()}] {fmt % args}")

    def _send(self, code, payload, content_type="application/json"):
        body = (json.dumps(payload, ensure_ascii=False).encode("utf-8")
                if content_type == "application/json" else payload)
        self.send_response(code)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Access-Control-Allow-Origin", "*")
        # Allow the service worker to be served from any path scope
        if content_type == "application/javascript":
            self.send_header("Service-Worker-Allowed", "/")
        self.end_headers()
        self.wfile.write(body)

    def _serve_static(self, path):
        if path in ("/", ""):
            path = "/index.html"
        fp = os.path.join(STATIC_DIR, path.lstrip("/"))
        if not os.path.abspath(fp).startswith(STATIC_DIR):
            self._send(403, {"error": "forbidden"}); return
        if not os.path.isfile(fp):
            self._send(404, {"error": "not found"}); return
        ctype = ("text/html" if fp.endswith(".html")
                 else "application/javascript" if fp.endswith(".js")
                 else "text/css" if fp.endswith(".css")
                 else "application/octet-stream")
        with open(fp, "rb") as fh:
            self._send(200, fh.read(), content_type=ctype)

    def do_GET(self):
        parsed = urllib.parse.urlparse(self.path)
        qs = urllib.parse.parse_qs(parsed.query)
        if parsed.path == "/api/geocode":
            q = qs.get("q", [""])[0]
            if not q:
                self._send(400, {"error": "missing q"}); return
            try:
                self._send(200, {"results": geocode(q)})
            except Exception as e:
                self._send(500, {"error": str(e)})
        elif parsed.path == "/api/reverse":
            try:
                lat = float(qs.get("lat", ["0"])[0])
                lon = float(qs.get("lon", ["0"])[0])
                self._send(200, reverse_geocode(lat, lon))
            except Exception as e:
                self._send(500, {"error": str(e)})
        elif parsed.path == "/api/vehicles":
            self._send(200, {"vehicles": VEHICLES, "max_shift_hours": MAX_SHIFT_HOURS})
        elif parsed.path == "/api/sessions":
            self._send(200, {"sessions": list_sessions()})
        elif parsed.path == "/api/session":
            fn = qs.get("file", [""])[0]
            try:
                self._send(200, load_session(fn))
            except FileNotFoundError:
                self._send(404, {"error": "session not found"})
            except Exception as e:
                self._send(500, {"error": str(e)})
        else:
            self._serve_static(parsed.path)

    def do_POST(self):
        parsed = urllib.parse.urlparse(self.path)
        length = int(self.headers.get("Content-Length", 0))
        raw = self.rfile.read(length) if length else b"{}"
        try:
            data = json.loads(raw.decode("utf-8"))
        except Exception:
            self._send(400, {"error": "invalid json"}); return

        if parsed.path == "/api/route":
            try:
                stops = data["stops"]
                fuel = data.get("fuel", {})
                vehicle = data.get("vehicle", {})
                optimize = bool(data.get("optimize", True))
                roundtrip = bool(data.get("roundtrip", False))
                max_hours = data.get("max_hours")
                speed_kmh = data.get("speed_kmh")
                plan = build_route_plan(stops, fuel, vehicle, optimize,
                                        roundtrip, max_hours, speed_kmh)
                paths = write_documents(plan)
                plan["saved_files"] = paths
                plan["saved_dir"] = ROUTES_DIR
                self._send(200, plan)
            except Exception as e:
                self._send(500, {"error": str(e)})
        elif parsed.path == "/api/session/save":
            try:
                name = data.get("name", "session")
                self._send(200, save_session(name, data.get("data", {})))
            except Exception as e:
                self._send(500, {"error": str(e)})
        else:
            self._send(404, {"error": "unknown endpoint"})


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--port", type=int, default=8777)
    ap.add_argument("--host", default="127.0.0.1")
    args = ap.parse_args()
    srv = ThreadingHTTPServer((args.host, args.port), Handler)
    print(f"CTS Route Planner running at http://{args.host}:{args.port}")
    print(f"Routes  -> {ROUTES_DIR}")
    print(f"Sessions-> {SESSIONS_DIR}")
    print(f"XLSX support: {HAVE_XLSX}")
    try:
        srv.serve_forever()
    except KeyboardInterrupt:
        print("\nShutting down.")
        srv.shutdown()


if __name__ == "__main__":
    main()
