#!/usr/bin/env python3
"""
Full self-test for the CTS Route Planner.
Exercises every endpoint + feature and asserts expected behaviour.
Run server first:  python server.py --port 8790
Then:              python self_test.py
"""
import json, os, glob, urllib.request, urllib.parse, sys

BASE = "http://127.0.0.1:8790"
ROUTES = os.path.join(os.path.dirname(os.path.abspath(__file__)), "routes")
SESSIONS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "sessions")

PASS, FAIL = [], []
def check(name, cond, detail=""):
    (PASS if cond else FAIL).append(name)
    print(f"  [{'PASS' if cond else 'FAIL'}] {name}" + (f"  — {detail}" if detail and not cond else ""))

def get(path):
    with urllib.request.urlopen(BASE+path, timeout=40) as r:
        return r.status, json.loads(r.read().decode())

def post(path, payload):
    req = urllib.request.Request(BASE+path, data=json.dumps(payload).encode(),
                                 headers={"Content-Type":"application/json"})
    with urllib.request.urlopen(req, timeout=90) as r:
        return r.status, json.loads(r.read().decode())

def raw(path):
    with urllib.request.urlopen(BASE+path, timeout=20) as r:
        return r.status, r.read(), dict(r.headers)

STOPS = [
    {"name":"Tuen Mun Hospital","nickname":"Depot","lat":22.404423,"lon":113.968754,"service_min":20},
    {"name":"Tuen Mun Police","nickname":"Police","lat":22.3866519,"lon":113.968301,"service_min":15},
    {"name":"TW point","nickname":"TW","lat":22.36892,"lon":114.1129885,"service_min":15},
    {"name":"Tai Hing Estate","nickname":"大興邨","lat":22.402933,"lon":113.9702707,"service_min":15},
    {"name":"Butterfly Estate","nickname":"蝴蝶邨","lat":22.3762852,"lon":113.9620179,"service_min":15},
    {"name":"Leung King Estate","nickname":"良景邨","lat":22.4062251,"lon":113.9619449,"service_min":15},
    {"name":"Siu Hei Court","nickname":"兆禧苑","lat":22.3738017,"lon":113.9667096,"service_min":15},
    {"name":"Melody Garden","nickname":"美樂花園","lat":22.3743067,"lon":113.9607311,"service_min":15},
]
FUEL = {"price":19.5,"economy":20,"unit":"l_per_100km"}
VEH = {"key":"truck_5t","label":"5-tonne truck"}

print("\n=== 1. Static assets ===")
for p, token in [("/","CTS Route Planner"),("/app.js","generate"),("/app.css",":root"),("/sw.js","TILE_CACHE")]:
    try:
        s, body, hdr = raw(p)
        check(f"GET {p} 200", s==200, f"status {s}")
        check(f"{p} content", token.encode() in body)
        if p=="/sw.js":
            check("sw.js Service-Worker-Allowed header", hdr.get("Service-Worker-Allowed")=="/", str(hdr.get("Service-Worker-Allowed")))
    except Exception as e:
        check(f"GET {p}", False, str(e))

print("\n=== 2. Geocode / reverse ===")
try:
    s, d = get("/api/geocode?"+urllib.parse.urlencode({"q":"Tuen Mun Hospital"}))
    check("geocode returns results", d.get("results") and len(d["results"])>0)
    check("geocode has lat/lon", "lat" in d["results"][0] and "lon" in d["results"][0])
except Exception as e: check("geocode", False, str(e))
try:
    s, d = get("/api/reverse?lat=22.404423&lon=113.968754")
    check("reverse returns name", bool(d.get("name")))
except Exception as e: check("reverse", False, str(e))
# edge: empty query
try:
    s, d = get("/api/geocode?q=")
    check("geocode empty q -> 400", d.get("error") is not None)
except urllib.error.HTTPError as e:
    check("geocode empty q -> 400", e.code==400)
except Exception as e: check("geocode empty q", False, str(e))

print("\n=== 3. Vehicles / params presets ===")
try:
    s, d = get("/api/vehicles")
    v = d.get("vehicles",{})
    check("3 vehicle presets", len(v)>=3)
    check("truck_5t has speed_kmh", v.get("truck_5t",{}).get("speed_kmh")==35.0)
    check("max_shift_hours present", d.get("max_shift_hours")==7.0)
except Exception as e: check("vehicles", False, str(e))

print("\n=== 4. Route: basic optimize ===")
try:
    s, plan = post("/api/route", {"stops":STOPS,"fuel":FUEL,"vehicle":VEH,"optimize":True,"roundtrip":False})
    check("route 200 no error", "error" not in plan, plan.get("error",""))
    check("8 stops in breakdown", len(plan["breakdown"])==8)
    check("optimized order differs from input", plan["order_indices"]!=list(range(8)))
    check("first stop is depot (index 0 fixed)", plan["order_indices"][0]==0)
    check("totals cost > 0", plan["totals"]["cost"]>0)
    check("geometry present", len(plan.get("geometry",[]))>10)
    check("matrix 8x8", len(plan["matrix"]["distances_km"])==8 and len(plan["matrix"]["distances_km"][0])==8)
    check("nickname used as label", plan["breakdown"][0]["name"]=="Depot")
    check("chinese nickname preserved", any("大興邨"==b["name"] for b in plan["breakdown"]))
    check("saved_files has xlsx", "xlsx" in plan.get("saved_files",{}))
    check("saved_files has csv/md/json/matrix", all(k in plan["saved_files"] for k in ["csv","markdown","json","matrix_csv"]))
    base_cost = plan["totals"]["cost"]
    base_drive = plan["totals"]["drive_min"]
except Exception as e:
    check("route basic", False, str(e)); base_cost=base_drive=0

print("\n=== 5. Param: speed override ===")
try:
    s, p0 = post("/api/route", {"stops":STOPS,"fuel":FUEL,"vehicle":VEH,"optimize":True,"speed_kmh":0})
    s, p35 = post("/api/route", {"stops":STOPS,"fuel":FUEL,"vehicle":VEH,"optimize":True,"speed_kmh":35})
    check("speed 0 uses OSRM (drive differs from 35)", p0["totals"]["drive_min"]!=p35["totals"]["drive_min"],
          f"{p0['totals']['drive_min']} vs {p35['totals']['drive_min']}")
    # at 35km/h, drive_min should equal distance/speed
    expect = p35["totals"]["distance_km"]/35*60
    check("35km/h drive matches dist/speed", abs(p35["totals"]["drive_min"]-expect)<1.0,
          f"{p35['totals']['drive_min']} vs {expect:.1f}")
    check("speed echoed in fuel.speed_kmh", p35["fuel"]["speed_kmh"]==35.0)
except Exception as e: check("speed override", False, str(e))

print("\n=== 6. Param: max_hours limit flip ===")
try:
    s, hi = post("/api/route", {"stops":STOPS,"fuel":FUEL,"vehicle":VEH,"optimize":True,"speed_kmh":35,"max_hours":7})
    s, lo = post("/api/route", {"stops":STOPS,"fuel":FUEL,"vehicle":VEH,"optimize":True,"speed_kmh":35,"max_hours":1})
    check("7h -> within limit", hi["shift"]["within_limit"] is True)
    check("1h -> over limit", lo["shift"]["within_limit"] is False)
    check("over_by_min > 0 when over", lo["shift"]["over_by_min"]>0)
    check("max_hours echoed", lo["shift"]["max_hours"]==1.0)
except Exception as e: check("max_hours", False, str(e))

print("\n=== 7. Param: fuel price/economy scaling ===")
try:
    s, a = post("/api/route", {"stops":STOPS,"fuel":{"price":10,"economy":20,"unit":"l_per_100km"},"vehicle":VEH,"optimize":True,"speed_kmh":35})
    s, b = post("/api/route", {"stops":STOPS,"fuel":{"price":20,"economy":20,"unit":"l_per_100km"},"vehicle":VEH,"optimize":True,"speed_kmh":35})
    check("double price ~ double cost", abs(b["totals"]["cost"]-2*a["totals"]["cost"])<0.1,
          f"{a['totals']['cost']} -> {b['totals']['cost']}")
    s, c = post("/api/route", {"stops":STOPS,"fuel":{"price":10,"economy":40,"unit":"l_per_100km"},"vehicle":VEH,"optimize":True,"speed_kmh":35})
    check("double economy ~ double cost (L/100km)", abs(c["totals"]["cost"]-2*a["totals"]["cost"])<0.1,
          f"{a['totals']['cost']} -> {c['totals']['cost']}")
    # km/L unit
    s, d = post("/api/route", {"stops":STOPS,"fuel":{"price":10,"economy":5,"unit":"km_per_l"},"vehicle":VEH,"optimize":True,"speed_kmh":35})
    check("km_per_l unit computes cost", d["totals"]["cost"]>0)
except Exception as e: check("fuel scaling", False, str(e))

print("\n=== 8. Param: service time affects delivery time ===")
try:
    s0 = [dict(x, service_min=0) for x in STOPS]
    s30 = [dict(x, service_min=30) for x in STOPS]
    _, r0 = post("/api/route", {"stops":s0,"fuel":FUEL,"vehicle":VEH,"optimize":True,"speed_kmh":35})
    _, r30 = post("/api/route", {"stops":s30,"fuel":FUEL,"vehicle":VEH,"optimize":True,"speed_kmh":35})
    check("service 0 -> service_min total 0", r0["totals"]["service_min"]==0)
    check("service 30x8 -> 240 min", abs(r30["totals"]["service_min"]-240)<0.1, str(r30["totals"]["service_min"]))
    check("more service -> more delivery time", r30["totals"]["delivery_min"]>r0["totals"]["delivery_min"])
except Exception as e: check("service time", False, str(e))

print("\n=== 9. Roundtrip + no-optimize ===")
try:
    _, rt = post("/api/route", {"stops":STOPS,"fuel":FUEL,"vehicle":VEH,"optimize":True,"roundtrip":True,"speed_kmh":35})
    _, ow = post("/api/route", {"stops":STOPS,"fuel":FUEL,"vehicle":VEH,"optimize":True,"roundtrip":False,"speed_kmh":35})
    check("roundtrip distance >= one-way", rt["totals"]["distance_km"]>=ow["totals"]["distance_km"])
    _, no = post("/api/route", {"stops":STOPS,"fuel":FUEL,"vehicle":VEH,"optimize":False,"speed_kmh":35})
    check("no-optimize keeps input order", no["order_indices"]==list(range(8)))
except Exception as e: check("roundtrip/no-opt", False, str(e))

print("\n=== 10. Error handling ===")
try:
    try:
        s, d = post("/api/route", {"stops":[STOPS[0]],"fuel":FUEL})
        check("1 stop -> error", d.get("error") is not None)
    except urllib.error.HTTPError as e:
        check("1 stop -> error", e.code==500)
    try:
        s, d = post("/api/route", {"stops":[],"fuel":FUEL})
        check("0 stops -> error", d.get("error") is not None)
    except urllib.error.HTTPError as e:
        check("0 stops -> error", e.code==500)
except Exception as e: check("error handling", False, str(e))

print("\n=== 11. Sessions save/list/load ===")
try:
    sess_data = {"vehicle":"truck_5t","fuel":FUEL,"max_hours":7,"speed_kmh":35,
                 "default_service":15,"optimize":True,"roundtrip":False,
                 "stops":[{"name":"A","nickname":"大興邨","lat":22.4,"lon":113.97,"service_min":20}],
                 "center":{"lat":22.39,"lng":113.98},"zoom":13}
    _, sv = post("/api/session/save", {"name":"selftest 中文","data":sess_data})
    check("session saved", "saved" in sv)
    _, ls = get("/api/sessions")
    found = [x for x in ls["sessions"] if x["file"]==sv["saved"]]
    check("session in list", len(found)==1)
    check("session name preserved", found[0]["name"]=="selftest 中文")
    _, lo = get("/api/session?file="+urllib.parse.quote(sv["saved"]))
    check("load round-trips nickname", lo["data"]["stops"][0]["nickname"]=="大興邨")
    check("load round-trips params", lo["data"]["speed_kmh"]==35 and lo["data"]["max_hours"]==7)
    # cleanup
    os.remove(os.path.join(SESSIONS, sv["saved"]))
    check("session cleanup", not os.path.exists(os.path.join(SESSIONS, sv["saved"])))
    # missing session
    try:
        s, d = get("/api/session?file=nonexistent_xyz.json")
        check("missing session -> 404/error", d.get("error") is not None)
    except urllib.error.HTTPError as e:
        check("missing session -> 404/error", e.code==404)
except Exception as e: check("sessions", False, str(e))

print("\n=== 12. Output files on disk ===")
try:
    import csv as _csv
    from openpyxl import load_workbook
    xs = sorted([f for f in glob.glob(ROUTES+"/*.xlsx") if not os.path.basename(f).startswith("~$")], key=os.path.getmtime)
    check("xlsx written", len(xs)>0)
    wb = load_workbook(xs[-1])
    check("xlsx 4 sheets", set(["Summary","Per-stop breakdown","Distance matrix km","Duration matrix min"]).issubset(set(wb.sheetnames)))
    cs = sorted(glob.glob(ROUTES+"/*_breakdown.csv"), key=os.path.getmtime)
    with open(cs[-1],"rb") as fh: bom = fh.read(3)==b"\xef\xbb\xbf"
    check("csv has utf-8 BOM", bom)
    with open(cs[-1], encoding="utf-8-sig") as fh:
        rows=list(_csv.reader(fh))
    check("csv chinese decodes", any(any("\u4e00"<=ch<="\u9fff" for ch in r[1]) for r in rows[1:]))
except Exception as e: check("output files", False, str(e))

print("\n" + "="*54)
print(f"RESULT: {len(PASS)} passed, {len(FAIL)} failed")
if FAIL:
    print("FAILURES:", ", ".join(FAIL)); sys.exit(1)
print("ALL TESTS PASSED")
